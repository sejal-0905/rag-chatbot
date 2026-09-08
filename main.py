import os
import re
import uuid
import asyncio
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor

from dotenv import load_dotenv
load_dotenv()

from fastapi import FastAPI, UploadFile, File
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Optional, List
import json

from pypdf import PdfReader
import pdfplumber
import docx

from groq import Groq
from tavily import TavilyClient
import chromadb
from sentence_transformers import SentenceTransformer, CrossEncoder
import nltk
try:
    nltk.data.find('tokenizers/punkt_tab')
except LookupError:
    nltk.download('punkt_tab', quiet=True)

CRISIS_PATTERNS = [
    r"\bkill (myself|me)\b",
    r"\bsuicid(e|al|ing)\b",
    r"\b(end(ing)?|take|taking) (my|this) (own )?life\b",
    r"\bwant(ing)? to die\b",
    r"\b(don'?t|do not) want to (be alive|live|exist)\b",
    r"\bno (reason|point)(\s+(to|in))?\s+liv(e|ing)\b",
    r"\bself[\s-]?harm(ing)?\b",
    r"\bhurt(ing)? myself\b",
    r"\bcut(ting)? myself\b",
    r"\bnot worth living\b",
    r"\b(overdose|od)\s+(on|of)\b",
    r"\bbetter off dead\b",
]
CRISIS_REGEX = re.compile("|".join(CRISIS_PATTERNS), re.IGNORECASE)

CRISIS_RESPONSE = (
    "I'm really glad you said something, but this isn't something I can help with here. "
    "If you're having thoughts of suicide or self-harm, please reach out to someone who can "
    "actually support you right now:\n\n"
    "- **Tele-MANAS**: Call **14416** (or **1-800-891-4416**) — free, government-run, 24/7, available in English and regional languages\n"
    "- **AASRA**: Call **9820466726** — 24/7 crisis helpline\n"
    "- **KIRAN Mental Health Helpline**: Call **1800-599-0019** — free, government-run, 24/7\n"
    "- **Outside India**: [findahelpline.com](https://findahelpline.com) lists crisis lines by country\n\n"
    "If you or someone else is in immediate danger, please contact emergency services on **112** right away.\n\n"
    "Please consider reaching out to one of these resources, or to someone you trust."
)

CATEGORY_RESPONSES = {
    "hate_speech": "I'm not able to help with that. I'm here to help you with your uploaded documents and related research topics — happy to help with that instead.",
    "violence": "I'm not able to help with that. I'm here to help you with your uploaded documents and related research topics — happy to help with that instead.",
    "illegal": "I'm not able to help with that. I'm here to help you with your uploaded documents and related research topics — happy to help with that instead.",
    "off_topic": "I'm built for document Q&A and research, not general conversation. Try asking me something about your uploaded documents or a related topic you're studying!",
}


def is_crisis_message(text: str) -> bool:
    return bool(CRISIS_REGEX.search(text or ""))


def is_doc_targeted_query(question: str, known_sources: list) -> bool:
    q = (question or "").strip().lower()
    if not q.startswith("from "):
        return False
    return any(src.lower() in q for src in known_sources)


DOC_RELEVANCE_THRESHOLD = 0.0
LOW_CONFIDENCE_THRESHOLD = 3.0

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

embedding_model = SentenceTransformer("BAAI/bge-base-en-v1.5")
reranker = CrossEncoder("cross-encoder/ms-marco-MiniLM-L-6-v2")

groq_client = Groq(api_key=os.environ.get("GROQ_API_KEY"))
tavily_client = TavilyClient(api_key=os.environ.get("TAVILY_API_KEY"))

chroma_client = chromadb.PersistentClient(path="./chroma_db")
collection = chroma_client.get_or_create_collection(
    name="documents",
    metadata={"hnsw:space": "cosine"}
)

executor = ThreadPoolExecutor(max_workers=4)
_sources_cache: list = []

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}


def invalidate_sources_cache():
    global _sources_cache
    _sources_cache = []


def embed_batch(texts: list):
    return embedding_model.encode(
        texts,
        normalize_embeddings=True,
        batch_size=32,
        show_progress_bar=False,
    ).tolist()


def semantic_chunk(text: str, max_chunk_size: int = 600, overlap_sentences: int = 1) -> list:
    sentences = nltk.sent_tokenize(text)
    chunks = []
    current_chunk = []
    current_size = 0
    for sentence in sentences:
        sentence_len = len(sentence)
        if current_size + sentence_len > max_chunk_size and current_chunk:
            chunks.append(" ".join(current_chunk))
            current_chunk = current_chunk[-overlap_sentences:] if overlap_sentences > 0 else []
            current_size = sum(len(s) for s in current_chunk)
        current_chunk.append(sentence)
        current_size += sentence_len
    if current_chunk:
        chunks.append(" ".join(current_chunk))
    return [c for c in chunks if c.strip()]


def extract_text_from_pdf(content: bytes) -> str:
    text = ""
    try:
        with pdfplumber.open(BytesIO(content)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + " "
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        for row in table:
                            row_text = " | ".join(cell or "" for cell in row)
                            text += row_text + " "
    except Exception:
        pass
    if not text.strip():
        try:
            reader = PdfReader(BytesIO(content))
            text = " ".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            pass
    return text.strip()


def extract_text_from_docx(content: bytes) -> str:
    try:
        doc = docx.Document(BytesIO(content))
        parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    parts.append(row_text)

        return " ".join(parts).strip()
    except Exception:
        return ""


def extract_text(content: bytes, filename: str) -> str:
    ext = "." + filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == ".pdf":
        return extract_text_from_pdf(content)
    elif ext == ".docx":
        return extract_text_from_docx(content)
    else:
        return content.decode("utf-8", errors="ignore")


def rewrite_query(question: str) -> list:
    try:
        response = groq_client.chat.completions.create(
            model="llama-3.1-8b-instant",
            messages=[{
                "role": "user",
                "content": (
                    "Generate 2 different ways to phrase this question for document search.\n"
                    "Return ONLY a JSON array of 2 strings, no explanation, no markdown, no extra text.\n"
                    'Example: ["phrase 1", "phrase 2"]\n'
                    f"Question: {question}"
                )
            }],
            max_tokens=100,
            temperature=0.3
        )
        text = response.choices[0].message.content.strip()
        text = text.replace("```json", "").replace("```", "").strip()
        variants = json.loads(text)
        if isinstance(variants, list):
            return [question] + variants[:2]
        return [question]
    except Exception:
        return [question]

CLASSIFIER_CATEGORIES = {"safe", "crisis", "hate_speech", "violence", "illegal", "off_topic"}

SAFETY_POLICY = """You are classifying user messages sent to a document Q&A chatbot.

Classify the message into exactly one category:

- crisis: any sign of suicidal thoughts, self-harm, or hopelessness about being alive — including indirect phrasing like "what's the point of being alive" or "I don't see a future for myself", not just direct statements.
- hate_speech: hateful, discriminatory, or harassing content targeting a person or group, OR requests for creative writing, roleplay, speeches, or fiction that promotes, glorifies, or normalizes hate speech, ethnic violence, or genocide against a real or fictional group.
- violence: requests for help planning, carrying out, or getting instructions related to violence or weapons, OR requests to praise, glorify, justify, or write favorably about real perpetrators or acts of violence/terrorism.
- illegal: requests for help with illegal activity (drugs, hacking, fraud, etc).
- off_topic: pure conversational chit-chat with no informational intent — jokes, small talk, creative writing requests, roleplay, or trivial requests like "what's 2+2". This is NOT for genuine questions seeking real information.
- safe: any genuine informational, factual, or research question, whether or not it's specifically about the uploaded documents.

Respond with ONLY a JSON object in this exact format, nothing else:
{"category": "<one of: safe, crisis, hate_speech, violence, illegal, off_topic>"}
"""

OUTPUT_SAFETY_POLICY = """You are reviewing an AI-generated answer that was already given to a user, to check if it's appropriate to show.

Classify the answer into exactly one category:

- crisis: the answer encourages, instructs, or fails to redirect away from suicide or self-harm.
- hate_speech: the answer contains hateful, discriminatory, or harassing content.
- violence: the answer provides instructions or assistance for violence or weapons.
- illegal: the answer provides instructions or assistance for illegal activity.
- safe: the answer is a normal, appropriate response — including factual answers, weather reports, biographical information, summaries, explanations, or any other helpful content. This is the default for ordinary, harmless answers.

Respond with ONLY a JSON object in this exact format, nothing else:
{"category": "<one of: safe, crisis, hate_speech, violence, illegal>"}
"""

def classify_question(question: str) -> str:
    try:
        response = groq_client.chat.completions.create(
            model="openai/gpt-oss-safeguard-20b",
            messages=[
                {"role": "system", "content": SAFETY_POLICY},
                {"role": "user", "content": question},
            ],
            max_tokens=200,
            temperature=0,
        )
        raw = (response.choices[0].message.content or "").strip()
        if not raw:
            return "classifier_error"
        parsed = clean_and_parse_json(raw)
        category = str(parsed.get("category", "")).strip().lower()
        return category if category in CLASSIFIER_CATEGORIES else "classifier_error"
    except Exception:
        return "classifier_error"
    
def classify_output(answer: str) -> str:
    try:
        response = groq_client.chat.completions.create(
            model="openai/gpt-oss-safeguard-20b",
            messages=[
                {"role": "system", "content": OUTPUT_SAFETY_POLICY},
                {"role": "user", "content": answer},
            ],
            max_tokens=200,
            temperature=0,
        )
        raw = (response.choices[0].message.content or "").strip()
        if not raw:
            return "safe"
        parsed = clean_and_parse_json(raw)
        category = str(parsed.get("category", "")).strip().lower()
        valid_output_categories = {"safe", "crisis", "hate_speech", "violence", "illegal"}
        return category if category in valid_output_categories else "safe"
    except Exception:
        return "safe"


def resolve_query(question: str, history: list) -> str:
    TRIGGERS = [
        "it", "its", "this", "that", "they", "them", "these", "those",
        "he", "she", "the same", "the above", "mentioned", "said", "such",
        "one", "ones",
    ]
    q_lower = question.lower().strip()
    needs_context = any(re.search(rf"\b{re.escape(t)}\b", q_lower) for t in TRIGGERS)
    if not needs_context or not history:
        return question
    recent_user_msgs = [m["content"] for m in history if m["role"] == "user"][-2:]
    context_str = " ".join(recent_user_msgs)
    return f"{context_str} {question}".strip()


def get_all_sources() -> list:
    global _sources_cache
    if _sources_cache:
        return _sources_cache
    try:
        result = collection.get()
        if result["metadatas"]:
            _sources_cache = sorted(set(m["source"] for m in result["metadatas"]))
            return _sources_cache
        return []
    except Exception:
        return []


def _do_retrieval(search_query: str, all_embs: list, col_count: int, n_candidates: int):
    all_candidates = {}
    for variant_emb in all_embs:
        try:
            res = collection.query(
                query_embeddings=[variant_emb],
                n_results=min(10, col_count),
                include=["documents", "metadatas", "distances"]
            )
            for chunk, meta, dist in zip(
                res["documents"][0], res["metadatas"][0], res["distances"][0]
            ):
                if chunk not in all_candidates or dist < all_candidates[chunk][1]:
                    all_candidates[chunk] = (meta, dist)
        except Exception:
            pass

    candidates = sorted(
        [(chunk, meta, dist) for chunk, (meta, dist) in all_candidates.items()],
        key=lambda x: x[2]
    )[:n_candidates]

    pairs = [[search_query, chunk] for chunk, _, _ in candidates]
    rerank_scores = reranker.predict(pairs, batch_size=32, show_progress_bar=False).tolist()
    return candidates, rerank_scores


async def retrieve_and_rerank(body, n_candidates: int = 15, n_top: int = 5):
    loop = asyncio.get_event_loop()
    history_dicts = [{"role": m.role, "content": m.content} for m in (body.history or [])]
    search_query = resolve_query(body.question, history_dicts)

    variants_future = loop.run_in_executor(executor, rewrite_query, search_query)
    original_emb_future = loop.run_in_executor(executor, embed_batch, [search_query])
    query_variants, original_embs = await asyncio.gather(variants_future, original_emb_future)

    remaining_variants = [v for v in query_variants if v != search_query]
    if remaining_variants:
        remaining_embs = await loop.run_in_executor(executor, embed_batch, remaining_variants)
    else:
        remaining_embs = []

    all_embs = original_embs + remaining_embs
    col_count = collection.count()

    candidates, rerank_scores = await loop.run_in_executor(
        executor, _do_retrieval, search_query, all_embs, col_count, n_candidates
    )

    ranked = sorted(zip(rerank_scores, candidates), key=lambda x: x[0], reverse=True)
    top = ranked[:n_top]
    top_chunks = [c[0] for _, c in top]

    top_debug = [
        {
            "rank": i + 1,
            "score": round(float(score), 4),
            "source": meta["source"],
            "chunk": chunk,
        }
        for i, (score, (chunk, meta, _)) in enumerate(top)
    ]

    source_counts: dict = {}
    for score, c in top:
        src = c[1]["source"]
        source_counts[src] = source_counts.get(src, 0) + max(0.0, float(score))

    if source_counts:
        max_s = max(source_counts.values())
        sources = [s for s, v in source_counts.items() if max_s == 0 or v >= max_s * 0.5]
        total = sum(source_counts.values())
        source_percentages = {
            s: round((v / total) * 100) if total > 0 else 0
            for s, v in sorted(source_counts.items(), key=lambda x: x[1], reverse=True)
        }
    else:
        sources = []
        source_percentages = {}

    return top_chunks, top_debug, sources, source_percentages, search_query


def tavily_search(query: str):
    try:
        response = tavily_client.search(query=query, search_depth="basic", max_results=5)
        results = response.get("results", [])
        web_context = ""
        web_sources = []
        for r in results:
            web_context += f"Title: {r.get('title')}\nContent: {r.get('content')}\nURL: {r.get('url')}\n\n"
            web_sources.append({"title": r.get("title", ""), "url": r.get("url", "")})
        return web_context.strip(), web_sources
    except Exception:
        return "", []


def clean_and_parse_json(raw: str):
    raw = raw.replace("```json", "").replace("```", "").strip()
    first = next((i for i, c in enumerate(raw) if c in "{["), -1)
    last = next((i for i in range(len(raw)-1, -1, -1) if raw[i] in "}]"), -1)
    if first != -1 and last != -1 and last > first:
        raw = raw[first:last+1]
    parsed = json.loads(raw)
    if isinstance(parsed, str):
        parsed = json.loads(parsed)
    return parsed


@app.get("/")
def home():
    return FileResponse("index.html")


@app.get("/health")
def health():
    return {"message": "RAG backend is running"}


@app.get("/documents")
def documents():
    try:
        results = collection.get()
        sources = (
            list(set(m["source"] for m in results["metadatas"]))
            if results["metadatas"] else []
        )
        return {"documents": sources}
    except Exception as e:
        return {"documents": [], "error": str(e)}


@app.get("/documents/chunks")
def document_chunks():
    try:
        results = collection.get()
        chunk_counts: dict = {}
        for m in results["metadatas"]:
            src = m["source"]
            chunk_counts[src] = chunk_counts.get(src, 0) + 1
        return chunk_counts
    except Exception as e:
        return {}


@app.delete("/documents/{filename}")
def delete_document(filename: str):
    try:
        existing = collection.get(where={"source": filename})
        if not existing["ids"]:
            return {"error": "Document not found"}
        collection.delete(ids=existing["ids"])
        invalidate_sources_cache()
        return {"deleted": filename}
    except Exception as e:
        return {"error": str(e)}


@app.post("/upload")
async def upload(file: UploadFile = File(...)):
    async def run():
        try:
            ext = "." + file.filename.lower().rsplit(".", 1)[-1] if "." in file.filename else ""
            if ext not in SUPPORTED_EXTENSIONS:
                yield json.dumps({
                    "error": f"Unsupported file type '{ext}'. Supported: PDF, DOCX, TXT, MD"
                }) + "\n"
                return

            content = await file.read()
            text = extract_text(content, file.filename)
            text = " ".join(text.split())

            if not text.strip():
                yield json.dumps({"error": "No text could be extracted from the file."}) + "\n"
                return

            chunks = semantic_chunk(text)
            if not chunks:
                yield json.dumps({"error": "No chunks produced from file."}) + "\n"
                return

            yield json.dumps({"progress": f"Extracted {len(chunks)} chunks, embedding..."}) + "\n"
            loop = asyncio.get_event_loop()
            embeddings = await loop.run_in_executor(executor, embed_batch, chunks)
            yield json.dumps({"progress": "Storing in database..."}) + "\n"

            try:
                existing = collection.get(where={"source": file.filename})
                if existing["ids"]:
                    collection.delete(ids=existing["ids"])
            except Exception:
                pass

            collection.add(
                ids=[str(uuid.uuid4()) for _ in chunks],
                embeddings=embeddings,
                documents=chunks,
                metadatas=[{"source": file.filename} for _ in chunks]
            )
            invalidate_sources_cache()
            yield json.dumps({"filename": file.filename, "chunks_stored": len(chunks)}) + "\n"
        except Exception as e:
            yield json.dumps({"error": str(e), "filename": file.filename, "chunks_stored": 0}) + "\n"

    return StreamingResponse(run(), media_type="application/x-ndjson")


class Message(BaseModel):
    role: str
    content: str


class AskRequest(BaseModel):
    question: str
    history: Optional[List[Message]] = []


class StudyRequest(BaseModel):
    mode: str
    context: str
    answer: Optional[str] = ""


STUDY_PROMPTS = {
    "notes": (
        "Based on the context below, create concise study notes.\n"
        "Return ONLY a JSON object with keys 'title' (string) and 'points' (array of strings).\n"
        "No markdown, no extra text, no explanation. Just the JSON.\n\n"
        "Context:\n{context}"
    ),
    "flashcards": (
        "Based on the context below, create 4 flashcards.\n"
        "Return ONLY a JSON array of objects with keys 'front' and 'back'.\n"
        "No markdown, no extra text, no explanation. Just the JSON.\n\n"
        "Context:\n{context}"
    ),
    "quiz": (
        "Based on the context below, create 3 multiple choice questions.\n"
        "Return ONLY a JSON array of objects with keys 'question', 'options' (object with keys a/b/c/d), 'answer' (one of a/b/c/d), 'explanation'.\n"
        "No markdown, no extra text, no explanation. Just the JSON.\n\n"
        "Context:\n{context}"
    ),
}


@app.post("/study")
async def study(body: StudyRequest):
    try:
        if body.mode not in STUDY_PROMPTS:
            return {"error": f"Unknown mode: {body.mode}"}
        prompt = STUDY_PROMPTS[body.mode].format(context=body.context[:4000])
        loop = asyncio.get_event_loop()
        response = await loop.run_in_executor(
            executor,
            lambda: groq_client.chat.completions.create(
                model="openai/gpt-oss-120b",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=1500,
                temperature=0.1,
            )
        )
        raw = response.choices[0].message.content.strip()
        try:
            parsed = clean_and_parse_json(raw)
        except json.JSONDecodeError:
            return {"error": f"Model returned malformed JSON: {raw[:200]}"}
        return {"result": parsed}
    except Exception as e:
        return {"error": str(e)}
    

def fixed_response_stream(message: str) -> StreamingResponse:
    async def _stream():
        yield "data: " + json.dumps({
            "sources": [], "source_percentages": {}, "debug_chunks": [],
            "web_used": False, "web_sources": [],
        }) + "\n\n"
        yield "data: " + json.dumps({"token": message}) + "\n\n"
        yield "data: " + json.dumps({"done": True}) + "\n\n"
    return StreamingResponse(_stream(), media_type="text/event-stream")

@app.post("/ask")
async def ask(body: AskRequest):
    if is_crisis_message(body.question):
        return fixed_response_stream(CRISIS_RESPONSE)

    loop = asyncio.get_event_loop()
    category = await loop.run_in_executor(executor, classify_question, body.question)

    if category == "crisis":
        return fixed_response_stream(CRISIS_RESPONSE)
    if category in CATEGORY_RESPONSES:
        return fixed_response_stream(CATEGORY_RESPONSES[category])

    try:
        if collection.count() == 0:
            async def empty_stream():
                yield "data: " + json.dumps({"token": "No documents uploaded yet."}) + "\n\n"
                yield "data: " + json.dumps({"done": True, "sources": []}) + "\n\n"
            return StreamingResponse(empty_stream(), media_type="text/event-stream")

        top_chunks, top_debug, doc_sources, source_percentages, search_query = await retrieve_and_rerank(body)

        all_sources = get_all_sources()
        docs_list = ", ".join(all_sources) if all_sources else "none"

        best_doc_score = max((c["score"] for c in top_debug), default=float("-inf"))
        doc_targeted = is_doc_targeted_query(body.question, all_sources)
        docs_are_relevant = doc_targeted or best_doc_score > DOC_RELEVANCE_THRESHOLD
        low_confidence = docs_are_relevant and not doc_targeted and best_doc_score < LOW_CONFIDENCE_THRESHOLD

        if docs_are_relevant:
            web_context, web_sources = "", []
            context = "\n\n".join(top_chunks)
            context_label = "Document context"
        else:
            web_context, web_sources = await loop.run_in_executor(executor, tavily_search, body.question)
            if web_context:
                context = web_context
                context_label = "Web context"
            else:
                context = "\n\n".join(top_chunks)
                context_label = "Document context"

        system_msg = {
            "role": "system",
            "content": (
                f"You are a helpful assistant. "
                f"Documents available in the knowledge base: {docs_list}.\n\n"
                "Answer using only the context provided below. It is labeled as either document "
                "context (from the uploaded files) or web context (from a live search) - never claim "
                "information came from the uploaded documents unless it actually appears in document context.\n"
                "Use markdown formatting where it helps clarity.\n"
                "The user may use synonyms or alternative terms — connect these to what's in the context and answer directly.\n"
                "For follow-up questions using 'it', 'this', 'that', resolve the reference from conversation history first.\n\n"
                + ("If the context above only weakly relates to the question, say so clearly rather than answering with full confidence.\n\n" if low_confidence else "")
                + f"{context_label}:\n" + context
            )
        }

        history_msgs = [{"role": m.role, "content": m.content} for m in (body.history or [])][-6:]
        messages = [system_msg] + history_msgs + [{"role": "user", "content": body.question}]

        async def stream_response():
            try:
                yield "data: " + json.dumps({
                    "sources": doc_sources,
                    "source_percentages": source_percentages,
                    "debug_chunks": top_debug,
                    "web_used": len(web_sources) > 0,
                    "web_sources": web_sources,
                    "low_confidence": low_confidence,
                }) + "\n\n"

                loop = asyncio.get_event_loop()
                stream = await loop.run_in_executor(
                    executor,
                    lambda: groq_client.chat.completions.create(
                        model="openai/gpt-oss-120b",
                        messages=messages,
                        max_tokens=2048,
                        temperature=0.2,
                        stream=True
                    )
                )
                full_answer = ""
                for chunk in stream:
                    token = chunk.choices[0].delta.content
                    if token:
                        full_answer += token
                        yield "data: " + json.dumps({"token": token}) + "\n\n"

                output_category = await loop.run_in_executor(executor, classify_output, full_answer)
                if output_category != "safe":
                    yield "data: " + json.dumps({
                        "token": "\n\n*This response was flagged after generation and may not be appropriate. Please rephrase your question.*"
                    }) + "\n\n"

                yield "data: " + json.dumps({"done": True}) + "\n\n"
            except Exception as e:
                yield "data: " + json.dumps({"error": str(e)}) + "\n\n"

        return StreamingResponse(stream_response(), media_type="text/event-stream")

    except Exception as e:
        err_msg = str(e)
        async def error_stream():
            yield "data: " + json.dumps({"error": err_msg}) + "\n\n"
        return StreamingResponse(error_stream(), media_type="text/event-stream")